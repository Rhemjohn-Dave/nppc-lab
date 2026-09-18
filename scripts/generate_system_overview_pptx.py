"""Generate NPPC Laboratory Management System overview PowerPoint."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# NPPC brand blues from the app
NAVY = RGBColor(0x1A, 0x36, 0x94)
BLUE = RGBColor(0x36, 0x5B, 0xB0)
LIGHT_BLUE = RGBColor(0x52, 0x82, 0xD3)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF4, 0xF7, 0xFB)
ACCENT_GREEN = RGBColor(0x05, 0x96, 0x69)

OUT = Path(__file__).resolve().parent.parent / "docs" / "NPPC_Lab_System_Overview.pptx"


def set_run(run, text, size=18, bold=False, color=SLATE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_top_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()


def add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(10), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "NPPC Laboratory Management System  ·  Confidential", 10, False, MUTED)
    num = slide.shapes.add_textbox(Inches(11.5), Inches(7.1), Inches(1.3), Inches(0.3))
    np_ = num.text_frame.paragraphs[0]
    np_.alignment = PP_ALIGN.RIGHT
    run = np_.add_run()
    set_run(run, f"{page} / {total}", 10, False, MUTED)


def title_block(slide, title: str, subtitle: str | None = None):
    add_top_bar(slide)
    t = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.6))
    p = t.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, title, 28, True, NAVY)
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(12), Inches(0.4))
        p = s.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, subtitle, 14, False, MUTED)


def bullet_slide(prs, title, bullets, subtitle=None, page=1, total=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, title, subtitle)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11.8), Inches(5.3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        set_run(run, f"•  {line}", 16, False, SLATE)
    add_footer(slide, page, total)
    return slide


def card(slide, left, top, width, height, heading, lines, fill=WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    shape.line.width = Pt(1)

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, Inches(0.1), height
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = NAVY
    accent.line.fill.background()

    h = slide.shapes.add_textbox(
        left + Inches(0.25), top + Inches(0.15), width - Inches(0.4), Inches(0.4)
    )
    p = h.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, heading, 16, True, NAVY)

    body = slide.shapes.add_textbox(
        left + Inches(0.25), top + Inches(0.55), width - Inches(0.4), height - Inches(0.7)
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        set_run(run, f"• {line}", 12, False, SLATE)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 12

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.8), Inches(13.333), Inches(1.7)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(1))
    p = t.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "NPPC Laboratory Management System", 36, True, WHITE)

    s = slide.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(11.5), Inches(0.8))
    p = s.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(
        run,
        "End-to-end lab workflow: intake → costing → JO approval → analysis → result release",
        18,
        False,
        RGBColor(0xD6, 0xE0, 0xF5),
    )

    f = slide.shapes.add_textbox(Inches(0.8), Inches(6.15), Inches(11), Inches(0.8))
    p = f.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "System overview  ·  Roles & functions  ·  September 2026", 14, False, WHITE)
    np_ = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.3), Inches(0.3))
    p = np_.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    set_run(run, f"1 / {total}", 10, False, WHITE)

    # 2. Agenda
    bullet_slide(
        prs,
        "Agenda",
        [
            "What the system is and what it solves",
            "End-to-end laboratory workflow",
            "Four main roles: Admin, Receiving, Analyst, Head Analysis",
            "Core modules and daily functions",
            "Documents: Job Order / RFA vs result forms",
            "Key operating rules (assignment, print, release)",
        ],
        page=2,
        total=total,
    )

    # 3. What it is
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, "What is this system?", "Laboratory Information Management for NPPC")
    points = [
        ("Receive", "Customer requests via public intake kiosk"),
        ("Track", "Samples, packages, and test progress"),
        ("Encode", "Analysts enter results on a shared work queue"),
        ("Control", "Official RFA and result PDF overlays"),
        ("Release", "Head releases results; ready for pickup"),
        ("Trace", "History, notifications, and print audit"),
    ]
    for i, (h, d) in enumerate(points):
        col = i % 3
        row = i // 3
        left = Inches(0.6 + col * 4.15)
        top = Inches(1.55 + row * 2.4)
        card(slide, left, top, Inches(3.9), Inches(2.15), h, [d])
    add_footer(slide, 3, total)

    # 4. Workflow
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, "End-to-end workflow", "From customer request to wet-signed papers")
    steps = [
        ("1", "Intake", "Customer submits samples & tests"),
        ("2", "Price", "Receiving costs the Job Order"),
        ("3", "Approve JO", "Head approves costing"),
        ("4", "Print & send", "Receiving prints JO ×3, sends to analysts"),
        ("5", "Encode", "Analysts enter results"),
        ("6", "Send to Head", "Analyst name/PRC, then submit"),
        ("7", "Release", "Head releases results"),
        ("8", "Packet", "Print result + reprint JO; wet-sign"),
    ]
    for i, (n, h, d) in enumerate(steps):
        col = i % 4
        row = i // 4
        left = Inches(0.45 + col * 3.2)
        top = Inches(1.55 + row * 2.5)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.0), Inches(2.15)
        )
        shape.adjustments[0] = 0.1
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.15), top + Inches(0.2), Inches(0.45), Inches(0.45)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = NAVY
        badge.line.fill.background()
        bt = slide.shapes.add_textbox(
            left + Inches(0.15), top + Inches(0.25), Inches(0.45), Inches(0.4)
        )
        bp = bt.text_frame.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        run = bp.add_run()
        set_run(run, n, 14, True, WHITE)
        ht = slide.shapes.add_textbox(
            left + Inches(0.7), top + Inches(0.25), Inches(2.1), Inches(0.4)
        )
        run = ht.text_frame.paragraphs[0].add_run()
        set_run(run, h, 15, True, NAVY)
        dt = slide.shapes.add_textbox(
            left + Inches(0.2), top + Inches(0.85), Inches(2.6), Inches(1.1)
        )
        dt.text_frame.word_wrap = True
        run = dt.text_frame.paragraphs[0].add_run()
        set_run(run, d, 13, False, SLATE)
    add_footer(slide, 4, total)

    # 5. Roles overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, "Main user roles", "Each role has a focused workspace")
    roles = [
        ("Admin", "Setup & control", ["Users & roles", "Prices & packages", "Assignments", "Controlled forms"]),
        ("Receiving", "Bridge to the lab", ["Costing / pricing", "Print JO ×3", "Send to analysts", "Reprint after release"]),
        ("Analyst", "Encode tests", ["Work queue", "Draft & complete", "Send to Head", "Print dated results"]),
        ("Head Analysis", "Two checkpoints", ["Approve JO", "Review results", "Release / return", "History"]),
    ]
    for i, (name, tag, lines) in enumerate(roles):
        left = Inches(0.4 + i * 3.2)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.5), Inches(3.05), Inches(5.2)
        )
        shape.adjustments[0] = 0.06
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        hdr = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, Inches(1.5), Inches(3.05), Inches(1.15)
        )
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = NAVY
        hdr.line.fill.background()
        nt = slide.shapes.add_textbox(left + Inches(0.15), Inches(1.65), Inches(2.75), Inches(0.45))
        run = nt.text_frame.paragraphs[0].add_run()
        set_run(run, name, 18, True, WHITE)
        tg = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.1), Inches(2.75), Inches(0.35))
        run = tg.text_frame.paragraphs[0].add_run()
        set_run(run, tag, 12, False, RGBColor(0xC7, 0xD2, 0xFE))
        body = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.85), Inches(2.65), Inches(3.5))
        tf = body.text_frame
        tf.word_wrap = True
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            run = p.add_run()
            set_run(run, f"•  {line}", 14, False, SLATE)
    add_footer(slide, 5, total)

    # 6. Admin
    bullet_slide(
        prs,
        "Admin — functions",
        [
            "Manage user accounts and role access",
            "Maintain procedures, prices, and the analysis catalog",
            "Configure analysis packages and designated package analysts",
            "Assignments matrix: which analysts may encode which tests",
            "Controlled forms: RFA (Job Order), package result sheets, standalone sheets",
            "History access, control numbers, print history, and audit",
            "Can step into operational areas when needed",
        ],
        "Setup and governance of the LMIS",
        page=6,
        total=total,
    )

    # 7. Receiving
    bullet_slide(
        prs,
        "Receiving — functions",
        [
            "Review submitted Job Orders from intake",
            "Enter prices and quantities (costing) → awaits Head JO approval",
            "After JO approval: print 3 JO/RFA copies (customer, accounting, Head file)",
            "Send to analysts — starts analysis; soft assignment by qualification & load",
            "After Head releases results: reprint JO for the customer packet",
            "Does not set Ready for pickup — that comes only from result release",
            "Payment / cashier settlement stays outside this system",
        ],
        "Bridge between intake and laboratory processing",
        page=7,
        total=total,
    )

    # 8. Analyst
    bullet_slide(
        prs,
        "Analyst — functions",
        [
            "See suggested work and other open tests they are qualified for (shared PCs)",
            "Enter results (including Passed/Failed where required); save drafts",
            "When all lines are complete: Send to Head is the primary queue action",
            "Enter analyst name and PRC at Send to Head (not the Head user)",
            "Preview the filled result form before sending",
            "After Head Release: print the dated result form and wet-sign with Head",
            "Package designated analyst consolidates and submits when applicable",
        ],
        "Encode tests and prepare the result packet",
        page=8,
        total=total,
    )

    # 9. Head
    bullet_slide(
        prs,
        "Head Analysis — functions",
        [
            "Checkpoint 1 — Approve Job Order after Receiving costing",
            "Unlocks Receiving JO print and Send to analysts",
            "Checkpoint 2 — Review encoded results and Release",
            "Release sets Ready for pickup, Report/Release dates, customer email",
            "Return selected lines for correction (before release only)",
            "Does not print RFA (Receiving does); does not enter analyst name/PRC",
            "Ink on paper remains the legal signature; system Release freezes dates",
        ],
        "Two distinct checkpoints — do not confuse Approve JO with Release results",
        page=9,
        total=total,
    )

    # 10. Modules
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, "Core modules", "Where work happens in the app")
    mods = [
        ("Public intake", "Kiosk for customer, samples, packages & tests"),
        ("Dashboard", "Role KPIs, attention items, next actions"),
        ("Receiving", "Price → print JO → send / reprint"),
        ("Analyst", "Test queue, encode, Send to Head, print"),
        ("Head Analysis", "JO approval + Results release"),
        ("History", "Searchable ready-for-pickup jobs"),
        ("Admin", "Users, catalog, packages, forms, audit"),
        ("Controlled forms", "Official PDF overlays (RFA & results)"),
    ]
    for i, (h, d) in enumerate(mods):
        col = i % 4
        row = i // 4
        left = Inches(0.4 + col * 3.2)
        top = Inches(1.55 + row * 2.5)
        card(slide, left, top, Inches(3.05), Inches(2.25), h, [d])
    add_footer(slide, 10, total)

    # 11. Documents & rules
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    title_block(slide, "Documents & operating rules", "What gets printed and when")
    card(
        slide,
        Inches(0.5),
        Inches(1.5),
        Inches(6.0),
        Inches(5.1),
        "Two customer papers",
        [
            "JO / RFA — Receiving prints after Head JO approval (×3)",
            "Result form — Analyst prints after Head result Release",
            "Ready for pickup = results released (not JO print)",
            "Wet signatures happen on paper after dates are frozen",
        ],
    )
    card(
        slide,
        Inches(6.8),
        Inches(1.5),
        Inches(6.0),
        Inches(5.1),
        "Assignment & signatories",
        [
            "Admin Assignments = who may do which test",
            "Send to analysts suggests least-busy qualified person",
            "Any qualified analyst may encode on a shared PC",
            "Analyst enters name/PRC at Send to Head",
            "Head only Approves JO and Releases results",
        ],
    )
    add_footer(slide, 11, total)

    # 12. Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    t = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, "Questions?", 40, True, WHITE)
    s = slide.shapes.add_textbox(Inches(1.5), Inches(3.6), Inches(10.3), Inches(1.2))
    p = s.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(
        run,
        "NPPC Laboratory Management System — controlled workflow from intake to released results.",
        16,
        False,
        RGBColor(0xD6, 0xE0, 0xF5),
    )
    f = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4))
    p = f.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, f"Thank you  ·  {total} / {total}", 12, False, WHITE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
